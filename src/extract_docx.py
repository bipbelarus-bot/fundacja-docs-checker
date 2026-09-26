from pathlib import Path
from docx import Document


def extract_docx_text(path: str | Path) -> str:
    path = Path(path)
    doc = Document(path)
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)
